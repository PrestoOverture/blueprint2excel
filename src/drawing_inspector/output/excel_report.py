"""Write an Excel inspection checklist from assembled drawing data."""

from os import PathLike

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from drawing_inspector.models import Dimension, DrawingData

_BASIC_INFO_FIELDS = (
    ("图号", "drawing_number"),
    ("名称", "part_name"),
    ("材料", "material"),
    ("比例", "scale"),
    ("日期", "date"),
    ("未注公差标准", "tolerance_standard"),
)


def write_excel_report(
    drawing_data: DrawingData, output_path: str | PathLike[str]
) -> None:
    """Write a four-sheet inspection checklist workbook."""
    workbook = Workbook()
    basic_info = workbook.active
    basic_info.title = "图纸基本信息"
    checklist = workbook.create_sheet("检查基准书")
    holes = workbook.create_sheet("孔位汇总")
    tolerance_details = workbook.create_sheet("公差明细")

    _write_basic_info(basic_info, drawing_data)
    _write_checklist(checklist, drawing_data)
    _write_holes(holes, drawing_data)
    _write_tolerance_details(tolerance_details, drawing_data)

    workbook.save(output_path)


def _write_basic_info(sheet: Worksheet, drawing_data: DrawingData) -> None:
    for label, field_name in _BASIC_INFO_FIELDS:
        sheet.append((label, getattr(drawing_data, field_name)))
    sheet.append(("技术要求",))
    for requirement in drawing_data.tech_requirements:
        sheet.append((requirement,))


def _write_checklist(sheet: Worksheet, drawing_data: DrawingData) -> None:
    sheet.append(
        (
            "序号",
            "气泡号",
            "检查项目",
            "标称值",
            "上偏差",
            "下偏差",
            "公差来源",
            "测量工具建议",
            "测量结果",
        )
    )
    for sequence, dimension in enumerate(drawing_data.dimensions, start=1):
        dimension.bubble_id = sequence
        sheet.append(
            (
                sequence,
                sequence,
                _inspection_item(dimension),
                dimension.nominal,
                dimension.upper_deviation,
                dimension.lower_deviation,
                _tolerance_source(dimension),
                _measurement_tool(dimension),
                None,
            )
        )


def _write_holes(sheet: Worksheet, drawing_data: DrawingData) -> None:
    sheet.append(
        (
            "标签",
            "类型",
            "孔径",
            "深度",
            "沉孔径",
            "沉孔深度",
            "螺纹规格",
            "底孔",
            "数量",
            "背面加工",
        )
    )
    for hole in drawing_data.hole_features:
        sheet.append(
            (
                hole.label,
                hole.hole_type,
                hole.primary_diameter,
                hole.depth,
                hole.counterbore_diameter,
                hole.counterbore_depth,
                hole.thread_spec,
                hole.pilot_hole,
                hole.quantity,
                "是" if hole.back_side else "否",
            )
        )


def _write_tolerance_details(
    sheet: Worksheet, drawing_data: DrawingData
) -> None:
    sheet.append(("序号", "标称值", "上偏差", "下偏差", "公差来源", "类型"))
    for dimension in drawing_data.dimensions:
        sheet.append(
            (
                dimension.bubble_id,
                dimension.nominal,
                dimension.upper_deviation,
                dimension.lower_deviation,
                _tolerance_source(dimension),
                dimension.dimension_type,
            )
        )


def _inspection_item(dimension: Dimension) -> str:
    nominal = _format_nominal(dimension.nominal)
    if dimension.dimension_type == "diameter":
        return f"⌀{nominal}"
    if dimension.dimension_type == "radius":
        return f"R{nominal}"
    if dimension.dimension_type == "angle":
        return f"{nominal}°"
    return nominal


def _format_nominal(nominal: float) -> str:
    if nominal.is_integer():
        return str(int(nominal))
    return str(nominal)


def _tolerance_source(dimension: Dimension) -> str:
    if dimension.tolerance_grade is not None:
        return f"{dimension.tolerance_grade}查表（未核实）"  # noqa: RUF001
    if dimension.upper_deviation is not None:
        return "图纸标注"
    return "无公差信息"


def _measurement_tool(dimension: Dimension) -> str:
    if dimension.dimension_type == "angle":
        return "角度尺"
    if dimension.nominal <= 200:
        return "游标卡尺"
    return "钢卷尺"
