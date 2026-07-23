from pathlib import Path

from openpyxl import load_workbook

from drawing_inspector.models import Dimension, DrawingData, HoleFeature
from drawing_inspector.output.excel_report import write_excel_report
from drawing_inspector.pipeline import assemble_drawing_data

FIXTURE = Path("tests/fixtures/HJMZ01-01-01.pdf")


def test_write_excel_report_exports_fixture_checklist(tmp_path: Path) -> None:
    drawing_data = assemble_drawing_data(FIXTURE)
    output_path = tmp_path / "inspection-checklist.xlsx"

    write_excel_report(drawing_data, output_path)

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [
        "图纸基本信息",
        "检查基准书",
        "孔位汇总",
        "公差明细",
    ]

    basic_info_values = {
        cell.value
        for row in workbook["图纸基本信息"].iter_rows()
        for cell in row
    }
    assert "HJMZ01-01-01" in basic_info_values
    assert "6061 合金" in basic_info_values

    checklist = workbook["检查基准书"]
    assert checklist.max_row == 19
    assert [checklist.cell(row=row, column=2).value for row in range(2, 20)] == list(
        range(1, 19)
    )
    tolerance_sources = {
        checklist.cell(row=row, column=7).value for row in range(2, 20)
    }
    assert "图纸标注" in tolerance_sources
    assert any(
        isinstance(source, str) and source.endswith("查表（未核实）")  # noqa: RUF001
        for source in tolerance_sources
    )

    holes = workbook["孔位汇总"]
    assert holes.max_row == 12

    tolerance_details = workbook["公差明细"]
    assert tolerance_details.max_row == 19
    assert [
        tolerance_details.cell(row=row, column=1).value for row in range(2, 20)
    ] == [checklist.cell(row=row, column=2).value for row in range(2, 20)]

    assert [dimension.bubble_id for dimension in drawing_data.dimensions] == list(
        range(1, 19)
    )


def test_write_excel_report_formats_inspection_rows_and_optional_values(
    tmp_path: Path,
) -> None:
    dimensions = [
        Dimension(20.0, 0.1, -0.1, None, "diameter", 0.0, 0.0, None),
        Dimension(19.75, None, None, None, "radius", 0.0, 0.0, None),
        Dimension(45.0, 1.0, -1.0, "GB/T 1804-m", "angle", 0.0, 0.0, None),
        Dimension(920.0, None, None, None, "linear", 0.0, 0.0, None),
    ]
    drawing_data = DrawingData(
        drawing_number="D-1",
        part_name="Part",
        material="Steel",
        scale="1:1",
        date="2026-07-23",
        tech_requirements=["去毛刺", "清洁"],
        dimensions=dimensions,
        hole_features=[
            HoleFeature(
                label="A",
                hole_type="through",
                primary_diameter=10.0,
                depth="贯穿",
                counterbore_diameter=None,
                counterbore_depth=None,
                thread_spec=None,
                pilot_hole=None,
                quantity=2,
                back_side=False,
            )
        ],
        tolerance_standard="GB/T 1804-m",
    )
    output_path = tmp_path / "formatted.xlsx"

    write_excel_report(drawing_data, output_path)

    workbook = load_workbook(output_path)
    checklist = workbook["检查基准书"]
    assert [checklist.cell(row=row, column=3).value for row in range(2, 6)] == [
        "⌀20",
        "R19.75",
        "45°",
        "920",
    ]
    assert [checklist.cell(row=row, column=7).value for row in range(2, 6)] == [
        "图纸标注",
        "无公差信息",
        "GB/T 1804-m查表（未核实）",  # noqa: RUF001
        "无公差信息",
    ]
    assert [checklist.cell(row=row, column=8).value for row in range(2, 6)] == [
        "游标卡尺",
        "游标卡尺",
        "角度尺",
        "钢卷尺",
    ]
    assert checklist.cell(row=3, column=5).value is None
    assert checklist.cell(row=3, column=6).value is None
    assert checklist.cell(row=2, column=4).value == 20
    assert all(checklist.cell(row=row, column=9).value is None for row in range(2, 6))

    holes = workbook["孔位汇总"]
    assert holes.cell(row=2, column=3).value == 10
    assert holes.cell(row=2, column=4).value == "贯穿"
    assert all(holes.cell(row=2, column=column).value is None for column in range(5, 9))
    assert holes.cell(row=2, column=10).value == "否"
